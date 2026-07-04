"""Тест модуля ingestion на моках: создаём временные .docx/.pdf/.png/.xlsx и
проверяем, что ingest() возвращает объект, валидный по schemas/document.schema.json.
png и pdf-сканы проверяются через реальный OCR (Tesseract), тесты пропускаются,
если Tesseract не установлен в системе — см. README.md модуля, раздел "OCR".
"""
import docx
import openpyxl
import pytesseract
import pytest
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas

from modules.ingestion.ingest import (
    _reconstruct_text_from_chars,
    _zero_width_glyph_ratio,
    ingest,
    ingest_folder,
)
from modules.rag_core.retrieve import retrieve
from schemas.validate_schema import validate


def _tesseract_available() -> bool:
    try:
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


requires_tesseract = pytest.mark.skipif(
    not _tesseract_available(),
    reason="Tesseract OCR не установлен в системе (см. modules/ingestion/README.md, раздел OCR)",
)


def _make_text_image(text: str, path):
    image = Image.new("RGB", (700, 150), color="white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=60)
    draw.text((20, 30), text, fill="black", font=font)
    image.save(str(path))


def test_ingest_docx_matches_schema(tmp_path):
    file_path = tmp_path / "sample.docx"
    document = docx.Document()
    document.add_heading("Вещественный состав хвостов", level=1)
    document.add_paragraph("Хвосты содержат тонковкрапленное золото, ассоциированное с сульфидами.")
    document.save(str(file_path))

    result = ingest(str(file_path))

    validate(result, "document")
    assert result["metadata"]["file_type"] == "docx"
    assert "тонковкрапленное" in result["full_text"]


def test_ingest_pdf_matches_schema(tmp_path):
    # Латиница: базовый шрифт Helvetica в reportlab не содержит глифов кириллицы,
    # это ограничение тестовой фикстуры, а не ingest() — pdfplumber одинаково
    # извлекает текст любого реального PDF независимо от языка.
    file_path = tmp_path / "sample.pdf"
    c = canvas.Canvas(str(file_path))
    c.drawString(100, 750, "Regrinding tailings increases sulfide liberation.")
    c.save()

    result = ingest(str(file_path))

    validate(result, "document")
    assert result["metadata"]["file_type"] == "pdf"
    assert "Regrinding" in result["full_text"]


def test_ingest_xlsx_matches_schema(tmp_path):
    file_path = tmp_path / "sample.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Хвосты"
    sheet.append(["Класс крупности, мкм", "Выход, %", "Содержание Ni, %"])
    sheet.append(["-125+71", "24.5", "0.46"])
    sheet.append(["-71+45", "18.2", "0.31"])
    workbook.save(str(file_path))

    result = ingest(str(file_path))

    validate(result, "document")
    assert result["metadata"]["file_type"] == "xlsx"
    assert result["metadata"]["extra"]["sheet_count"] == 1
    assert len(result["tables"]) == 1
    assert result["tables"][0]["caption"] == "Хвосты"
    assert result["tables"][0]["data"][0] == ["Класс крупности, мкм", "Выход, %", "Содержание Ni, %"]
    assert "24.5" in result["full_text"]


@requires_tesseract
def test_ingest_png_matches_schema_and_runs_ocr(tmp_path):
    file_path = tmp_path / "sample.png"
    _make_text_image("ZOLOTO TAILINGS", file_path)

    result = ingest(str(file_path))

    validate(result, "document")
    assert result["metadata"]["file_type"] == "png"
    assert result["metadata"]["ocr_confidence"] is not None
    assert 0.0 <= result["metadata"]["ocr_confidence"] <= 1.0
    assert "ZOLOTO" in result["full_text"].upper()


@requires_tesseract
def test_ingest_pdf_scan_triggers_ocr(tmp_path):
    # PDF-страница без текстового слоя, только вставленное изображение —
    # имитирует "фото учебника/статьи" из реальной задачи.
    image_path = tmp_path / "page_image.png"
    _make_text_image("SULFIDE OXIDATION", image_path)

    file_path = tmp_path / "scan.pdf"
    c = canvas.Canvas(str(file_path), pagesize=(700, 150))
    c.drawImage(str(image_path), 0, 0, width=700, height=150)
    c.save()

    result = ingest(str(file_path))

    validate(result, "document")
    assert result["metadata"]["ocr_confidence"] is not None
    assert "SULFIDE" in result["full_text"].upper()


def _make_sample_folder(folder):
    docx_path = folder / "report.docx"
    document = docx.Document()
    document.add_heading("Вещественный состав хвостов", level=1)
    document.add_paragraph("Хвосты содержат тонковкрапленное золото, ассоциированное с сульфидами.")
    document.save(str(docx_path))

    xlsx_path = folder / "tailings.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Хвосты"
    sheet.append(["Класс крупности, мкм", "Содержание Ni, %"])
    sheet.append(["-125+71", "0.46"])
    workbook.save(str(xlsx_path))

    # временный Excel-лок должен игнорироваться ingest_folder
    (folder / "~$tailings.xlsx").write_bytes(b"")
    # неподдерживаемый формат должен игнорироваться молча
    (folder / "notes.txt").write_text("не документ базы знаний", encoding="utf-8")


def test_ingest_folder_caches_results_between_runs(tmp_path):
    source_folder = tmp_path / "source"
    source_folder.mkdir()
    _make_sample_folder(source_folder)
    cache_dir = tmp_path / "cache"

    first_run = ingest_folder(str(source_folder), cache_dir=str(cache_dir))
    assert len(first_run) == 2  # docx + xlsx; ~$-лок и .txt пропущены

    cached_files = list(cache_dir.glob("*.json"))
    assert len(cached_files) == 2

    second_run = ingest_folder(str(source_folder), cache_dir=str(cache_dir))
    # doc_id — новый uuid4 при каждом реальном парсинге; если он совпал между
    # запусками, значит второй запуск действительно взял результат из кэша,
    # а не распарсил файлы заново.
    assert {d["doc_id"] for d in first_run} == {d["doc_id"] for d in second_run}

    for doc in second_run:
        validate(doc, "document")


def test_ingest_folder_output_feeds_rag_core_without_errors(tmp_path):
    source_folder = tmp_path / "source"
    source_folder.mkdir()
    _make_sample_folder(source_folder)
    cache_dir = tmp_path / "cache"

    documents = ingest_folder(str(source_folder), cache_dir=str(cache_dir))

    # Формулировка нарочно пересекается по словам с текстом в report.docx
    # ("тонковкрапленное золото... ассоциированное с сульфидами") — TF-IDF
    # ищет по точному совпадению словоформ, а не по смыслу (см. README rag_core).
    query = {
        "target_property": "тонковкрапленное золото, ассоциированное с сульфидами в хвостах",
        "constraints": {"materials": [], "budget": None, "equipment": [], "regulatory": []},
    }
    retrieval_result = retrieve(query, documents)

    validate(retrieval_result, "retrieval_result")
    assert len(retrieval_result["retrieved_chunks"]) > 0


def _make_char(text, x0, x1, top):
    return {"text": text, "x0": x0, "x1": x1, "top": top}


def test_zero_width_glyph_ratio_detects_broken_font_metrics():
    # Реальный случай: некоторые старые сканированные pdf (см. README, раздел
    # "Известные проблемы pdf") хранят битые метаданные ширины глифов — у
    # большинства символов x1-x0 ~ 0, из-за чего pdfplumber переставляет буквы
    # местами при обычной кластеризации в слова.
    broken_chars = [_make_char(c, 10.0, 10.0, 5.0) for c in "тест"]
    assert _zero_width_glyph_ratio(broken_chars) == 1.0

    normal_chars = [_make_char(c, float(i), float(i) + 5.0, 5.0) for i, c in enumerate("тест")]
    assert _zero_width_glyph_ratio(normal_chars) == 0.0

    assert _zero_width_glyph_ratio([]) == 0.0


def test_reconstruct_text_from_chars_preserves_stream_order_by_line():
    chars = [_make_char(c, float(i), float(i), 10.0) for i, c in enumerate("привет")]
    chars += [_make_char(c, float(i), float(i), 30.0) for i, c in enumerate("мир")]

    text = _reconstruct_text_from_chars(chars)

    assert text == "привет\nмир"
